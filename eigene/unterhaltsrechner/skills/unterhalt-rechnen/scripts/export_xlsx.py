"""Excel-Berechnungsblatt für den Unterhaltsrechner (openpyxl)."""
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EUR = '#,##0.00 "€"'
KOPF = Font(bold=True, color="FFFFFF")
KOPF_FUELL = PatternFill("solid", fgColor="2F4F6F")
WARN = PatternFill("solid", fgColor="FCE4D6")


def _f(x):
    return float(x) if isinstance(x, Decimal) else x


def _kopf(ws, zeile, werte):
    for i, w in enumerate(werte, 1):
        c = ws.cell(row=zeile, column=i, value=w)
        c.font, c.fill = KOPF, KOPF_FUELL
        c.alignment = Alignment(wrap_text=True, vertical="top")


def _breiten(ws, breiten):
    for i, b in enumerate(breiten, 1):
        ws.column_dimensions[get_column_letter(i)].width = b


def schreibe(r, pfad, abschnitte, de):
    """r: Ergebnis von berechne(); abschnitte, de: Hilfsfunktionen aus unterhalt.py."""
    wb = Workbook()
    ll = r["leitlinie"]
    namen = []
    for m in r["monate"]:
        for n in m["soll"]:
            if n not in namen:
                namen.append(n)

    # Übersicht
    ws = wb.active
    ws.title = "Übersicht"
    titel = "VORLÄUFIGE Unterhaltsberechnung (Testphase)" if r.get("vorlaeufig") else "Unterhaltsberechnung (Testphase)"
    ws["A1"] = titel
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Leitlinie: {ll['name']} (OLG {', '.join(ll['olg'])})"
    ws["A3"] = f"Quelle: {ll['quelle']} ({ll['quelle_art']})"
    ws["A4"] = f"Rechner v{r['version']}, eigene Ergänzung von digitalmann, nicht von Klotzkette. Keine Rechtsberatung; Ergebnis vor Verwendung nachprüfen."
    if r.get("vorlaeufig"):
        ws["A5"] = "Vorläufig: Einkommen beruht ganz oder teilweise auf BWA, Schätzung oder ungeprüfter Angabe."
        ws["A5"].fill = WARN
    z = 7
    _kopf(ws, z, ["Zeitraum", "Monate", "DT", "bereinigtes Einkommen", "Gruppe"] + namen)
    for a in abschnitte(r["monate"]):
        z += 1
        m = a["m"]
        zr = a["von"] if a["von"] == a["bis"] else f"{a['von']} bis {a['bis']}"
        werte = [zr, a["anzahl"], m["dt"], _f(m["pflichtiger"]["gesamt"]), m["kinder_info"].get("gruppe")]
        werte += [_f(m["soll"].get(n, 0)) for n in namen]
        for i, w in enumerate(werte, 1):
            c = ws.cell(row=z, column=i, value=w)
            if i == 4 or i > 5:
                c.number_format = EUR
    _breiten(ws, [22, 8, 7, 20, 8] + [16] * len(namen))
    ws.freeze_panes = "A8"

    # Monate
    wm = wb.create_sheet("Monate")
    kopf = ["Monat", "DT", "bereinigtes Einkommen", "Gruppe"] + [f"Soll {n}" for n in namen] + [f"gezahlt {n}" for n in namen]
    _kopf(wm, 1, kopf)
    for i, m in enumerate(r["monate"], 2):
        werte = [m["monat"], m["dt"], _f(m["pflichtiger"]["gesamt"]), m["kinder_info"].get("gruppe")]
        werte += [_f(m["soll"].get(n, 0)) for n in namen] + [_f(m["ist"].get(n, 0)) for n in namen]
        for j, w in enumerate(werte, 1):
            c = wm.cell(row=i, column=j, value=w)
            if j == 3 or j > 4:
                c.number_format = EUR
    letzte = len(r["monate"]) + 1
    sz = letzte + 1
    wm.cell(row=sz, column=1, value="Summe").font = Font(bold=True)
    for j in range(5, len(kopf) + 1):
        col = get_column_letter(j)
        c = wm.cell(row=sz, column=j, value=f"=SUM({col}2:{col}{letzte})")
        c.number_format, c.font = EUR, Font(bold=True)
    _breiten(wm, [10, 7, 20, 8] + [16] * (2 * len(namen)))
    wm.freeze_panes = "B2"

    # Rechenweg
    wr = wb.create_sheet("Rechenweg")
    wr.column_dimensions["A"].width = 140
    z = 1
    for a in abschnitte(r["monate"]):
        m = a["m"]
        zr = a["von"] if a["von"] == a["bis"] else f"{a['von']} bis {a['bis']}"
        wr.cell(row=z, column=1, value=f"{zr} (Düsseldorfer Tabelle {m['dt']})").font = Font(bold=True, size=12)
        z += 1
        bloecke = [("Einkommen Pflichtiger", m["pflichtiger"]["schritte"])]
        ki = m["kinder_info"]
        if ki:
            zeilen = [ki.get("eingruppierung", "")] + ki.get("herabstufung", [])
            if ki.get("mangelfall"):
                zeilen.append("Mangelfall: " + ki["mangelfall"])
            bloecke.append(("Eingruppierung", zeilen))
        for k in m["kinder"]:
            bloecke.append((f"{k['name']} (Rang {k['rang']}, {k['status']})", k["schritte"]))
        if m["ehegatte"]:
            bloecke.append((f"Ehegattenunterhalt ({m['ehegatte']['art']})", m["ehegatte"]["schritte"]))
        for titel_b, zeilen in bloecke:
            wr.cell(row=z, column=1, value=titel_b).font = Font(bold=True)
            z += 1
            for s in zeilen:
                wr.cell(row=z, column=1, value=de(s)).alignment = Alignment(wrap_text=True)
                z += 1
        z += 1

    # Rückstand
    wk = wb.create_sheet("Rückstand")
    kopf = ["Berechtigter", "Soll gesamt", "Soll ab Verzug", "gezahlt ab Verzug", "Rückstand", "davon Land (UVG)", "davon Berechtigter",
            "älter als 1 Jahr", "möglicherweise verjährt", "ausgeschlossen § 1585b"]
    _kopf(wk, 1, kopf)
    for i, (n, x) in enumerate(r["rueckstand"].items(), 2):
        werte = [n, x["soll"], x["soll_ab_verzug"], x["ist_ab_verzug"], x["differenz"], x["land_uvg"], x["davon_kind"],
                 x["alt_verwirkung"], x["alt_verjaehrung"], x["ausgeschlossen_1585b"]]
        for j, w in enumerate(werte, 1):
            c = wk.cell(row=i, column=j, value=_f(w))
            if j > 1:
                c.number_format = EUR
    _breiten(wk, [16] + [17] * 9)

    # Hinweise
    wh = wb.create_sheet("Hinweise")
    wh.column_dimensions["A"].width = 150
    _kopf(wh, 1, ["Hinweise und Annahmen"])
    p = r["parameter"]
    zeilen = list(r["hinweise"]) + [
        f"Leitlinie berufsbedingte Aufwendungen (Nr. {p['berufsbedingt']['nr']}): {p['berufsbedingt']['text']}",
        f"Leitlinie Fahrtkosten (Nr. {p['fahrtkosten']['nr']}): {p['fahrtkosten']['text']}",
        f"Leitlinie Erwerbstätigenbonus (Nr. {p['bonus']['nr']}): {p['bonus']['text']}",
        f"Leitlinie Ehegattenselbstbehalt (Nr. {p['selbstbehalt_ehegatte']['nr']}): {p['selbstbehalt_ehegatte']['text']}",
        f"Leitlinie Eingruppierung: {p['eingruppierung']}",
        "Nicht automatisiert: Wechselmodell, Altersvorsorge- und Krankenvorsorgeunterhalt, Dreiteilung mit neuem Ehegatten, "
        "Befristung und Herabsetzung (§ 1578b BGB), Verwirkung, Netto-Ermittlung aus Brutto, Steuern.",
    ]
    for i, h in enumerate(zeilen, 2):
        c = wh.cell(row=i, column=1, value=de(h))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if h.startswith("VORLÄUFIG"):
            c.fill = WARN

    wb.save(pfad)
    return pfad
