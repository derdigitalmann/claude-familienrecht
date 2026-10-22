#!/usr/bin/env python3
"""Regressionstests mit handgerechneten Erwartungswerten. Aufruf: python3 tests/test_unterhalt.py"""
import json
import sys
from decimal import Decimal as D
from pathlib import Path

HIER = Path(__file__).resolve().parent
sys.path.insert(0, str(HIER.parent / "scripts"))
import unterhalt as u  # noqa: E402

FEHLER = []


def pruefe(name, ist, soll):
    if D(str(ist)) != D(str(soll)):
        FEHLER.append(f"{name}: ist {ist}, soll {soll}")
    else:
        print(f"ok  {name}: {soll}")


def basis(**kw):
    e = {"leitlinie": "suedl", "zeitraum": {"von": "2026-01", "bis": "2026-01"},
         "pflichtiger": {"erwerbstaetig": True, "einkommen": [{"ab": "2026-01", "netto": 2000}]}, "kinder": []}
    e.update(kw)
    return e


def m0(e):
    return u.berechne(e)["monate"][0]


# 1. Beispiel DT 2026 Anm. C (Mangelfall, drei Kinder, eines privilegiert volljährig)
r = m0(json.load(open(HIER / "dt-anm-c.json")))
pruefe("DT Anm. C K1", r["soll"]["K1"], "107.60")
pruefe("DT Anm. C K2", r["soll"]["K2"], "105.02")
pruefe("DT Anm. C K3", r["soll"]["K3"], "87.38")

# 2. Ein Kind, Höherstufung um eine Gruppe: 3.000 € -> Gruppe 4, ein Berechtigter -> Gruppe 5; Kind 8 Jahre: 670 - 129,50
r = m0(basis(pflichtiger={"einkommen": [{"ab": "2026-01", "netto": 3000}]}, kinder=[{"name": "A", "geburtsdatum": "2018-05-05"}]))
pruefe("Höherstufung Gruppe", r["kinder_info"]["gruppe"], 5)
pruefe("Höherstufung Zahlbetrag", r["soll"]["A"], "540.50")

# 3. Altersstufenwechsel im Geburtsmonat (§ 1612a Abs. 3): Kind wird im März 2026 6 Jahre
e = basis(zeitraum={"von": "2026-02", "bis": "2026-03"}, kinder=[{"name": "B", "geburtsdatum": "2020-03-31"}],
          pflichtiger={"einkommen": [{"ab": "2026-01", "netto": 2000}]})
ms = u.berechne(e)["monate"]
# Gruppe 1 +1 = 2, aber Rest 2.000 - 381,50 unter Bedarfskontrollbetrag 1.750 -> zurück auf Gruppe 1
pruefe("Stufe Februar", ms[0]["soll"]["B"], "356.50")   # 486 - 129,50
pruefe("Stufe März", ms[1]["soll"]["B"], "428.50")      # 558 - 129,50

# 4. Pauschale Niedersachsen: 5 % von 4.000 = 200, gedeckelt 150; von 600 = 30, Mindestbetrag 50
x = u.bereinige({"netto": 4000, "berufsbedingt": {"modus": "pauschal"}}, u.leitlinie_aufloesen(u.lade_daten()[1], "niedersachsen"), "P", set())
pruefe("Pauschale max", x["berufsbedingt"], "150.00")
x = u.bereinige({"netto": 600, "berufsbedingt": {"modus": "pauschal"}}, u.leitlinie_aufloesen(u.lade_daten()[1], "niedersachsen"), "P", set())
pruefe("Pauschale min", x["berufsbedingt"], "50.00")
x = u.bereinige({"netto": 4000, "berufsbedingt": {"modus": "pauschal"}}, u.leitlinie_aufloesen(u.lade_daten()[1], "suedl"), "P", set())
pruefe("Pauschale SüdL ohne Deckel", x["berufsbedingt"], "200.00")

# 5. Fahrtkosten Koblenz: 40 km -> 30 x 14 + 10 x 7 = 490
x = u.bereinige({"netto": 3000, "berufsbedingt": {"modus": "konkret", "fahrt": {"entfernung_km": 40}}}, u.leitlinie_aufloesen(u.lade_daten()[1], "koblenz"), "P", set())
pruefe("Fahrtkosten Koblenz", x["berufsbedingt"], "490.00")
# Schleswig-Holstein, Beispiel der Leitlinie Nr. 10.2.2: 50 km -> 667,33
x = u.bereinige({"netto": 3000, "berufsbedingt": {"modus": "konkret", "fahrt": {"entfernung_km": 50}}}, u.leitlinie_aufloesen(u.lade_daten()[1], "schleswig"), "P", set())
pruefe("Fahrtkosten SH (Leitlinienbeispiel)", x["berufsbedingt"], "667.33")

# 6. Ehegattenunterhalt ohne Kinder, SüdL: P 3.000, B 1.000 (beide erwerbstätig)
#    (2.700 - 900) / 2 = 900 -> aufgerundet 900
e = basis(pflichtiger={"einkommen": [{"ab": "2026-01", "netto": 3000}]},
          ehegatte={"erwerbstaetig": True, "einkommen": [{"ab": "2026-01", "netto": 1000}]})
pruefe("Ehegatte SüdL", m0(e)["soll"]["Ehegatte"], 900)
# Brandenburg: Bonus vor Abzug Kredit: Erwerb 3.000, Kredit 300 -> 3.000 x 0,1 = 300 Bonus; P = 2.700 - 300 = 2.400; B = 900; (2.400 - 900)/2 = 750
e["leitlinie"] = "brandenburg"
e["pflichtiger"]["einkommen"][0]["abzuege"] = [{"bezeichnung": "Kredit", "betrag": 300}]
pruefe("Ehegatte Brandenburg Bonus vorab", m0(e)["soll"]["Ehegatte"], 750)
#    SüdL zum Vergleich: (2.700 x 0,9 = 2.430 - 900)/2 = 765
e["leitlinie"] = "suedl"
pruefe("Ehegatte SüdL mit Kredit", m0(e)["soll"]["Ehegatte"], 765)

# 7. Selbstbehalt Ehegatte, nicht erwerbstätiger Pflichtiger (Rente 2.000): DT 1.475, Brandenburg 1.600
e = basis(pflichtiger={"erwerbstaetig": False, "einkommen": [{"ab": "2026-01", "netto": 0, "sonstige_einkuenfte": 2000}]},
          ehegatte={"einkommen": [{"ab": "2026-01", "netto": 0}]})
pruefe("SB Ehegatte DT 1.475", m0(e)["soll"]["Ehegatte"], 525)
e["leitlinie"] = "brandenburg"
pruefe("SB Ehegatte Brandenburg 1.600", m0(e)["soll"]["Ehegatte"], 400)

# 8. Rundung: NRW kaufmännisch, SüdL aufrunden. P 3.001 ohne Bonusfragen: Rente 3.001, B 0 -> 1.500,50
e = basis(pflichtiger={"erwerbstaetig": False, "einkommen": [{"ab": "2026-01", "netto": 0, "sonstige_einkuenfte": 3001}]},
          ehegatte={"einkommen": [{"ab": "2026-01", "netto": 0}]})
pruefe("Rundung SüdL", m0(e)["soll"]["Ehegatte"], 1501)
e["leitlinie"] = "nrw"
pruefe("Rundung NRW", m0(e)["soll"]["Ehegatte"], 1501)  # 1.500,50 kaufmännisch -> 1.501
e["pflichtiger"]["einkommen"][0]["sonstige_einkuenfte"] = 3000.8   # 1.500,40
pruefe("Rundung NRW abwärts", m0(e)["soll"]["Ehegatte"], 1500)
e["leitlinie"] = "suedl"
pruefe("Rundung SüdL aufwärts", m0(e)["soll"]["Ehegatte"], 1501)

# 9. Volljähriges Kind in Ausbildung (nicht privilegiert), beide Eltern: 3.000 + 2.000 = 5.000 -> Gruppe 9: 1.061
#    offen 1.061 - 259 = 802; Quote (3.000-1.750)/((3.000-1.750)+(2.000-1.750)) = 1.250/1.500 -> 668,33
e = basis(pflichtiger={"einkommen": [{"ab": "2026-01", "netto": 3000}]},
          kinder=[{"name": "V", "geburtsdatum": "2006-01-10", "status_ab_18": "volljaehrig"}],
          anderer_elternteil={"einkommen": [{"ab": "2026-01", "netto": 2000}]})
pruefe("Volljährig Quote", m0(e)["soll"]["V"], "668.33")

# 10. Student mit eigenem Haushalt, Pflichtiger allein: 990 - 259 = 731
e = basis(pflichtiger={"einkommen": [{"ab": "2026-01", "netto": 4000}]},
          kinder=[{"name": "S", "geburtsdatum": "2005-01-10", "status_ab_18": "student_eigener_haushalt"}])
pruefe("Student", m0(e)["soll"]["S"], "731.00")

# 11. Saarbrücken verweist auf NRW
ll = u.leitlinie_aufloesen(u.lade_daten()[1], "saarbruecken")
pruefe("Saarbrücken Rundung wie NRW", int(ll["rundung"]["modus"] == "kaufmaennisch"), 1)

# 12. Leitlinie fehlt -> Fehler
try:
    u.berechne(basis(leitlinie=None))
    FEHLER.append("Leitlinie fehlt: kein Fehler ausgelöst")
except u.Fehler:
    print("ok  Leitlinie fehlt -> Fehler")

# 13. Rückstand 2025 mit DT 2025: Kind 3 Jahre, 2.000 €, Gruppe 2 scheitert am Bedarfskontrollbetrag -> Gruppe 1: 482 - 127,50 = 354,50; gezahlt 300
e = basis(zeitraum={"von": "2025-06", "bis": "2025-06"}, pflichtiger={"einkommen": [{"ab": "2025-01", "netto": 2000}]}, kinder=[{"name": "R", "geburtsdatum": "2022-01-01"}],
          zahlungen=[{"monat": "2025-06", "an": "R", "betrag": 300}])
pruefe("Rückstand 2025", u.berechne(e)["rueckstand"]["R"]["differenz"], "54.50")

# 14. Privilegiert volljährig, Pflichtiger allein: keine Höherstufung. 3.000 -> Gruppe 4: 803 - 259 = 544
e = basis(pflichtiger={"einkommen": [{"ab": "2026-01", "netto": 3000}]},
          kinder=[{"name": "P", "geburtsdatum": "2007-01-10", "status_ab_18": "privilegiert"}])
pruefe("Privilegiert ohne Höherstufung", m0(e)["soll"]["P"], "544.00")
# Mit minderjährigem Geschwister: Gruppe 4 -> Rest 3.000 - 512,50 - 544 = 1.943,50 < BKB 1.950 -> Gruppe 3
e["kinder"].append({"name": "M", "geburtsdatum": "2015-06-01"})
r = m0(e)
pruefe("Geschwister minderjährig", r["soll"]["M"], "484.50")   # 614 - 129,50
pruefe("Geschwister privilegiert", r["soll"]["P"], "509.00")   # 768 - 259

# 15. Selbstständiger mit drei Steuerbescheiden: 193.000 : 36 = 5.361,11 - 1.233,33 - 780 - 600 = 2.747,78
#     Gruppe 3 -> 3 Berechtigte -> 2 -> Bedarfskontrolle -> 1; Kinder 12 und 8 Jahre; Ehegatte durch Selbstbehalt auf 195 € begrenzt
sd = {"zeitraeume": [{"bezeichnung": "2023", "quelle": "steuerbescheid", "gewinn": 58000},
                     {"bezeichnung": "2024", "quelle": "steuerbescheid", "gewinn": 64000},
                     {"bezeichnung": "2025", "quelle": "steuerbescheid", "gewinn": 71000}],
      "steuern_jahr": 14800, "kranken_pflege_monat": 780, "altersvorsorge_monat": 600, "entnahmen_geprueft": True}
e = basis(zeitraum={"von": "2026-03", "bis": "2026-03"},
          pflichtiger={"einkommen": [{"ab": "2026-01", "selbstaendig": sd}]},
          kinder=[{"name": "E", "geburtsdatum": "2013-05-20"}, {"name": "P", "geburtsdatum": "2018-02-11"}],
          ehegatte={"erwerbstaetig": True, "einkommen": [{"ab": "2026-01", "netto": 1150, "berufsbedingt": {"modus": "pauschal"}}]})
r = u.berechne(e)
m = r["monate"][0]
pruefe("Selbstständig Netto", m["pflichtiger"]["gesamt"], "2747.78")
pruefe("Selbstständig Kind 12", m["soll"]["E"], "523.50")
pruefe("Selbstständig Kind 8", m["soll"]["P"], "428.50")
pruefe("Selbstständig Ehegatte", m["soll"]["Ehegatte"], 195)
pruefe("Selbstständig nicht vorläufig", int(r["vorlaeufig"]), 0)
# Nur BWA ohne Steuern: vorläufig
sd2 = {"zeitraeume": [{"bezeichnung": "BWA 01-08/2026", "quelle": "bwa", "gewinn": 46400, "monate": 8}], "kranken_pflege_monat": 780}
e["pflichtiger"]["einkommen"][0]["selbstaendig"] = sd2
r = u.berechne(e)
pruefe("BWA vorläufig", int(r["vorlaeufig"]), 1)
pruefe("BWA Steuerhinweis", int(any("Steuern auf den Gewinn fehlen" in h for h in r["hinweise"])), 1)

# 16. Verzug: Rückstand erst ab Monatserstem des Verzugsmonats (§ 1613 Abs. 1 S. 2 BGB)
e = basis(zeitraum={"von": "2026-01", "bis": "2026-04"}, stichtag="2026-10",
          pflichtiger={"einkommen": [{"ab": "2026-01", "netto": 2000, "netto_basis": "jahresschnitt"}]},
          kinder=[{"name": "R", "geburtsdatum": "2022-01-01", "verzug_ab": "2026-03-17"}])
r = u.berechne(e)["rueckstand"]["R"]
pruefe("Verzug Soll gesamt (4 x 356,50)", r["soll"], "1426.00")
pruefe("Verzug Soll ab 03/2026 (2 x 356,50)", r["soll_ab_verzug"], "713.00")
# 17. Unterhaltsvorschuss: Übergang auf das Land in Höhe der Leistung (§ 7 UVG); hier 230 € je Monat, nichts gezahlt
e["kinder"][0]["unterhaltsvorschuss"] = [{"von": "2026-01", "bis": "2026-12", "betrag": 230}]
r = u.berechne(e)["rueckstand"]["R"]
pruefe("UVG Land", r["land_uvg"], "460.00")
pruefe("UVG Kind", r["davon_kind"], "253.00")
# 18. Ausbildungsvergütung Minderjähriger: (800 - 100) / 2 = 350 angerechnet; 16 Jahre, 3.000 € -> Gruppe 5: 784 - 129,50 - 350 = 304,50
e = basis(pflichtiger={"einkommen": [{"ab": "2026-01", "netto": 3000, "netto_basis": "jahresschnitt"}]},
          kinder=[{"name": "A", "geburtsdatum": "2009-09-01", "ausbildungsverguetung": 800}])
pruefe("Azubi minderjährig hälftig", m0(e)["soll"]["A"], "304.50")
# 19. Mehrbedarf Kita 200 € anteilig: Kind 4 J., P 3.000 -> Gruppe 5: 584 - 129,50 = 454,50;
#     Quote (3.000 - 454,50 - 1.750) / ((795,50) + (2.000 - 1.750)) = 795,5/1.045,5 -> 152,18
e = basis(pflichtiger={"einkommen": [{"ab": "2026-01", "netto": 3000, "netto_basis": "jahresschnitt"}]},
          anderer_elternteil={"einkommen": [{"ab": "2026-01", "netto": 2000, "netto_basis": "jahresschnitt"}]},
          kinder=[{"name": "K", "geburtsdatum": "2021-06-01", "mehrbedarf": 200}])
pruefe("Mehrbedarf anteilig", m0(e)["soll"]["K"], "606.68")
# 20. § 1585b Abs. 3: nachehelich, Rechtshängigkeit 2026-03 -> Monate vor 2025-03 ausgeschlossen (hier 2025-01, 2025-02)
e = basis(zeitraum={"von": "2025-01", "bis": "2025-04"}, stichtag="2026-10",
          pflichtiger={"erwerbstaetig": False, "einkommen": [{"ab": "2025-01", "netto": 0, "sonstige_einkuenfte": 3000}]},
          ehegatte={"art": "nachehelich", "rechtskraft_scheidung": "2024-06", "verzug_ab": "2024-07", "rechtshaengig_ab": "2026-03",
                    "einkommen": [{"ab": "2025-01", "netto": 0}]})
r = u.berechne(e)["rueckstand"]["Ehegatte"]
pruefe("1585b ausgeschlossen (2 x 1.500)", r["ausgeschlossen_1585b"], "3000.00")
pruefe("1585b verbleibend (2 x 1.500)", r["differenz"], "3000.00")

# 21. Excel-Export (nur wenn openpyxl installiert)
try:
    import openpyxl  # noqa: F401
    import tempfile
    import export_xlsx
    e = basis(kinder=[{"name": "X", "geburtsdatum": "2018-05-05"}], zahlungen=[{"monat": "2026-01", "an": "X", "betrag": 100}])
    r = u.berechne(e)
    with tempfile.TemporaryDirectory() as td:
        pfad = Path(td) / "t.xlsx"
        export_xlsx.schreibe(r, pfad, u.abschnitte, u.de)
        wb = openpyxl.load_workbook(pfad)
        pruefe("Excel Blätter", len(wb.sheetnames), 5)
        pruefe("Excel Zahlbetrag", wb["Übersicht"].cell(row=8, column=6).value, r["monate"][0]["soll"]["X"])
except ImportError:
    print("--  Excel-Test übersprungen (openpyxl fehlt)")

# 22. Eingabeprüfung: Fehleingaben enden mit Fehler, nicht mit Absturz oder stillem Ergebnis
faelle = {
    "Datum deutsch": basis(kinder=[{"name": "K", "geburtsdatum": "01.01.2020"}]),
    "Name doppelt": basis(kinder=[{"name": "K", "geburtsdatum": "2020-01-01"}, {"name": "K", "geburtsdatum": "2015-01-01"}]),
    "Status Tippfehler": basis(kinder=[{"name": "K", "geburtsdatum": "2005-01-01", "status_ab_18": "privilegiet"}]),
    "von nach bis": basis(zeitraum={"von": "2026-05", "bis": "2026-01"}),
    "Zahlung unbekannt": basis(kinder=[{"name": "K", "geburtsdatum": "2020-01-01"}], zahlungen=[{"monat": "2026-01", "an": "X", "betrag": 1}]),
    "Betrag als Text": basis(pflichtiger={"einkommen": [{"ab": "2026-01", "netto": "3.000,00"}]}),
    "Gruppe 16": basis(kinder=[{"name": "K", "geburtsdatum": "2020-01-01"}], gruppe={"fest": 16}),
    "Monate 0": basis(pflichtiger={"einkommen": [{"ab": "2026-01", "selbstaendig": {"zeitraeume": [{"gewinn": 1, "monate": 0}]}}]}),
}
for name, e in faelle.items():
    try:
        u.berechne(e)
        FEHLER.append(f"Eingabeprüfung {name}: kein Fehler ausgelöst")
    except u.Fehler:
        print(f"ok  Eingabeprüfung {name} -> Fehler")

# 23. Altersstufe bei Geburt am Monatsersten (§§ 187 Abs. 2 S. 2, 188 Abs. 2 BGB): geboren 01.03.2020 -> 2. Stufe ab Februar 2026
e = basis(zeitraum={"von": "2026-02", "bis": "2026-02"}, gruppe={"korrektur": 0},
          kinder=[{"name": "K", "geburtsdatum": "2020-03-01"}])
pruefe("Stufenwechsel Monatserster", m0(e)["soll"]["K"], "428.50")
#     Volljährigkeit erst ab dem Geburtstagsmonat: geboren 01.03.2008 -> Februar 2026 noch 3. Stufe (Gruppe 1: 653 - 129,50)
e = basis(zeitraum={"von": "2026-02", "bis": "2026-02"}, gruppe={"korrektur": 0},
          kinder=[{"name": "K", "geburtsdatum": "2008-03-01", "status_ab_18": "privilegiert"}])
pruefe("Volljährigkeit Monatserster", m0(e)["soll"]["K"], "523.50")
# 24. Aufrundung nie über Leistungsfähigkeit: Rente 2.950,60, SB 1.475 -> verfügbar 1.475,60, Bedarf 1.475,30 -> 1.475
e = basis(pflichtiger={"erwerbstaetig": False, "einkommen": [{"ab": "2026-01", "netto": 0, "sonstige_einkuenfte": 2950.6}]},
          ehegatte={"einkommen": [{"ab": "2026-01", "netto": 0}]})
pruefe("Rundung begrenzt durch Selbstbehalt", m0(e)["soll"]["Ehegatte"], 1475)
# 25. Koblenz-Staffel abschaltbar: 45 km x 14 = 630
x = u.bereinige({"netto": 3000, "berufsbedingt": {"modus": "konkret", "fahrt": {"entfernung_km": 45, "staffel": False}}},
                u.leitlinie_aufloesen(u.lade_daten()[1], "koblenz"), "P", set())
pruefe("Koblenz ohne Staffel", x["berufsbedingt"], "630.00")
# 26. Mangelfall überschreitet die Masse nicht: 2.200, Kinder 3/9/13, Masse 750
e = basis(leitlinie="nrw", pflichtiger={"einkommen": [{"ab": "2026-01", "netto": 2200}]},
          kinder=[{"name": "A", "geburtsdatum": "2023-01-01"}, {"name": "B", "geburtsdatum": "2017-01-01"},
                  {"name": "C", "geburtsdatum": "2013-01-01"}])
r = m0(e)
pruefe("Mangelfall Summe <= Masse", int(sum(D(str(v)) for v in r["soll"].values()) <= D("750")), 1)

# 27. Leitlinienfassungen 2024/2025 für Rückstände
lls = u.lade_daten()[1]
try:
    u.berechne(basis(leitlinie="niedersachsen", zeitraum={"von": "2025-01", "bis": "2025-01"},
                     pflichtiger={"einkommen": [{"ab": "2025-01", "netto": 2000}]}))
    FEHLER.append("Niedersachsen 2025 ohne olg: kein Fehler ausgelöst")
except u.Fehler:
    print("ok  Niedersachsen 2025 ohne olg -> Fehler")
try:
    u.berechne(basis(leitlinie="niedersachsen", olg="Hamm"))
    FEHLER.append("olg passt nicht zur Leitlinie: kein Fehler ausgelöst")
except u.Fehler:
    print("ok  olg passt nicht zur Leitlinie -> Fehler")
pb = {"netto": 4000, "berufsbedingt": {"modus": "pauschal"}}
pruefe("Celle 2025 Pauschale ohne Deckel", u.bereinige(pb, u.leitlinie_fuer(lls, "niedersachsen", 2025, "Celle"), "P", set())["berufsbedingt"], "200.00")
pruefe("Niedersachsen 2026 Pauschale gedeckelt", u.bereinige(pb, u.leitlinie_fuer(lls, "niedersachsen", 2026, "Celle"), "P", set())["berufsbedingt"], "150.00")
pruefe("Braunschweig 2024 Mindestpauschale unter 500 €", u.bereinige({"netto": 400, "berufsbedingt": {"modus": "pauschal"}},
       u.leitlinie_fuer(lls, "niedersachsen", 2024, "braunschweig"), "P", set())["berufsbedingt"], "25.00")
# Oldenburg 2025: Ehegattenunterhalt auf 5 € gerundet; Rente 4.006 -> 2.003 -> 2.005; 2026 einheitlich aufgerundet -> 2.003
e = basis(leitlinie="niedersachsen", olg="Oldenburg", zeitraum={"von": "2025-12", "bis": "2026-01"},
          pflichtiger={"erwerbstaetig": False, "einkommen": [{"ab": "2025-01", "netto": 0, "sonstige_einkuenfte": 4006}]},
          ehegatte={"einkommen": [{"ab": "2025-01", "netto": 0}]})
ms = u.berechne(e)["monate"]
pruefe("Oldenburg 2025 Rundung 5 €", ms[0]["soll"]["Ehegatte"], 2005)
pruefe("Niedersachsen 2026 Rundung 1 €", ms[1]["soll"]["Ehegatte"], 2003)
# NRW 2024: Hamm kaufmännisch, Köln aufrunden; Rente 3.000,60 -> 1.500,30
e = basis(leitlinie="nrw", olg="Hamm", zeitraum={"von": "2024-06", "bis": "2024-06"},
          pflichtiger={"erwerbstaetig": False, "einkommen": [{"ab": "2024-01", "netto": 0, "sonstige_einkuenfte": 3000.6}]},
          ehegatte={"einkommen": [{"ab": "2024-01", "netto": 0}]})
pruefe("Hamm 2024 kaufmännisch", m0(e)["soll"]["Ehegatte"], 1500)
e["olg"] = "Köln"
pruefe("Köln 2024 aufrunden", m0(e)["soll"]["Ehegatte"], 1501)
pruefe("NRW 2025 einheitlich", int(u.leitlinie_fuer(lls, "nrw", 2025)["berufsbedingt"]["pauschale"] == "nur_fiktiv"), 1)
# Saarbrücken: 2024 eigene Pressemitteilung (keine Übernahme NRW), 2026 Verweis auf NRW
s24 = u.leitlinie_fuer(lls, "saarbruecken", 2024)
pruefe("Saarbrücken 2024 Pauschale DT A.3", u.bereinige(pb, s24, "P", set())["berufsbedingt"], "150.00")
pruefe("Saarbrücken 2024 ohne Verweis", int(not s24.get("verweis")), 1)
r = u.berechne(basis(leitlinie="saarbruecken", zeitraum={"von": "2024-12", "bis": "2025-01"},
                     pflichtiger={"einkommen": [{"ab": "2024-01", "netto": 2000}]}))
pruefe("Saarbrücken Hinweis UNSICHER", int(any("[UNSICHER]" in h for h in r["hinweise"])), 1)
pruefe("Fassungen ausgewiesen", len(r["fassungen"]), 2)

# 28. Altersvorsorgeunterhalt zweistufig (Bremer Tabelle), SüdL 2026, nachehelich, P 4.000 erwerbstätig, B ohne Einkommen
#     Stufe 1: 4.000 x 0,9 / 2 = 1.800; + 25 % = 2.250 x 18,6 % = 418,50
#     Stufe 2: (4.000 - 418,50) = 3.581,50 - Bonus 358,15 = 3.223,35 / 2 = 1.611,68 -> 1.612; AVU aufgerundet 419
def avu_fall(netto, av, art="nachehelich", **kw):
    eh = {"art": art, "rechtskraft_scheidung": "2025-06-01", "altersvorsorge": av, "verzug_ab": "2025-12-01"}
    return basis(pflichtiger={"einkommen": [{"ab": "2026-01", "netto": netto}]}, ehegatte=eh, **kw)


r = m0(avu_fall(4000, {"bremer_zuschlag_prozent": 25}))
pruefe("AVU Elementar Stufe 2", r["soll"]["Ehegatte"], 1612)
pruefe("AVU Betrag", r["soll"]["Ehegatte Altersvorsorge"], 419)
#     Tabellenzeilen: 1.800 -> 21 %: 2.178 x 18,6 % = 405,11; Stufe 2: 3.594,89 - 359,49 = 3.235,40 / 2 = 1.617,70 -> 1.618; AVU 406
r = m0(avu_fall(4000, {"bremer_tabelle": [{"bis": 1500, "prozent": 15}, {"bis": 2000, "prozent": 21}]}))
pruefe("AVU Tabelle Elementar", r["soll"]["Ehegatte"], 1618)
pruefe("AVU Tabelle Betrag", r["soll"]["Ehegatte Altersvorsorge"], 406)
#     Mindestbedarf: P 2.500 -> 1.125 < 1.200 -> kein AVU
r = m0(avu_fall(2500, {"bremer_zuschlag_prozent": 25}))
pruefe("AVU entfällt bei Mindestbedarf", r["soll"]["Ehegatte Altersvorsorge"], 0)
#     Leistungsfähigkeit: P 3.000: Stufe 1 1.350 -> 313,88; Stufe 2 (2.686,12 - 268,61) / 2 = 1.208,76 -> 1.209; frei 1.400 - 1.209 = 191
r = m0(avu_fall(3000, {"bremer_zuschlag_prozent": 25}))
pruefe("AVU Elementar vorrangig", r["soll"]["Ehegatte"], 1209)
pruefe("AVU begrenzt durch Leistungsfähigkeit", r["soll"]["Ehegatte Altersvorsorge"], 191)
try:
    u.berechne(avu_fall(4000, {"bremer_zuschlag_prozent": 25}, art="trennung"))
    FEHLER.append("AVU Trennung ohne ab: kein Fehler ausgelöst")
except u.Fehler:
    print("ok  AVU Trennung ohne Rechtshängigkeit -> Fehler")
e = avu_fall(4000, {"bremer_zuschlag_prozent": 25}, zahlungen=[{"monat": "2026-01", "an": "Ehegatte Altersvorsorge", "betrag": 100}])
pruefe("AVU Rückstand", u.berechne(e)["rueckstand"]["Ehegatte Altersvorsorge"]["differenz"], 319)

# 29. Dreiteilung mit neuem Ehegatten: P 4.000 (anrechenbar 3.600), Geschiedene 0, Neue 1.000 erwerbstätig (900)
#     Halbteilung 1.800; Dreiteilung (3.600 + 0 + 900) / 3 = 1.500 -> 1.500
def dt_fall(rang):
    e = basis(pflichtiger={"einkommen": [{"ab": "2026-01", "netto": 4000}],
                           "neuer_ehegatte": {"rang": rang, "erwerbstaetig": True, "einkommen": [{"ab": "2026-01", "netto": 1000}]}},
              ehegatte={"art": "nachehelich", "rechtskraft_scheidung": "2025-06-01"})
    return m0(e)["soll"]["Ehegatte"]


pruefe("Dreiteilung Gleichrang", dt_fall("gleich"), 1500)
pruefe("Dreiteilung neuer Ehegatte nachrangig", dt_fall("neuer_nachrangig"), 1800)
try:
    u.berechne(basis(pflichtiger={"einkommen": [{"ab": "2026-01", "netto": 4000}], "neuer_ehegatte": {"rang": "gleich"}},
                     ehegatte={"art": "trennung"}))
    FEHLER.append("Dreiteilung beim Trennungsunterhalt: kein Fehler ausgelöst")
except u.Fehler:
    print("ok  Dreiteilung beim Trennungsunterhalt -> Fehler")

if FEHLER:
    print("\nFEHLER:\n" + "\n".join(FEHLER))
    sys.exit(1)
print("\nalle Tests bestanden")
